from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.repositories.command_repository import LotteryExportJob


class LotteryExcelExporter:
    """Idempotent Excel mirror for committed lottery registrations.

    SQLite is the source of truth. A registration can be exported repeatedly
    without creating duplicate rows, which makes crash recovery safe when the
    worker succeeds in Excel but crashes before acknowledging the outbox job.
    """

    HEADERS = [
        "Telegram User ID",
        "Name and Family Name",
        "Company A Registration Email",
        "Company B Registration Email",
        "Telegram ID",
        "Registered At",
    ]

    def __init__(self, path: str) -> None:
        self._path = Path(path)

    def upsert(self, job: LotteryExportJob) -> None:
        self._path.parent.mkdir(parents=True, exist_ok=True)
        temp = self._path.with_suffix(self._path.suffix + ".tmp")
        workbook = None
        try:
            if self._path.exists():
                workbook = load_workbook(self._path)
                sheet = workbook.active
                if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
                    sheet.delete_rows(1)
            else:
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Lottery Registrations"
                sheet.append(self.HEADERS)

            values = [
                job.telegram_user_id,
                job.full_name,
                job.company_a_email,
                job.company_b_email,
                f"@{job.telegram_username}" if job.telegram_username else "",
                job.registered_at,
            ]

            existing_row = None
            for row in range(2, sheet.max_row + 1):
                if sheet.cell(row, 1).value == job.telegram_user_id:
                    existing_row = row
                    break

            if existing_row is None:
                sheet.append(values)
            else:
                for column, value in enumerate(values, start=1):
                    sheet.cell(existing_row, column, value)

            workbook.save(temp)
            temp.replace(self._path)
        finally:
            if workbook is not None:
                workbook.close()
            if temp.exists():
                temp.unlink(missing_ok=True)
