from pathlib import Path

import pytest
import pytest_asyncio

from app.core.exceptions import RepositoryError
from app.database.connection import Database
from app.features.commands.export import LotteryExcelExporter
from app.features.commands.export_worker import LotteryExportWorker
from app.repositories.command_repository import SQLiteCommandRepository


@pytest_asyncio.fixture
async def database(tmp_path: Path):
    database = Database(str(tmp_path / "integration.db"))
    await database.connect()
    try:
        yield database
    finally:
        await database.close()


async def _outbox_row(database: Database, telegram_user_id: int):
    cursor = await database.connection.execute(
        """SELECT telegram_user_id, attempts, last_error
           FROM lottery_export_outbox WHERE telegram_user_id = ?""",
        (telegram_user_id,),
    )
    return await cursor.fetchone()


@pytest.mark.asyncio
async def test_registration_and_outbox_are_committed_atomically(database: Database):
    repository = SQLiteCommandRepository(database)

    registration = await repository.create_registration(
        telegram_user_id=1001,
        full_name="Integration User",
        company_a_email="a@example.com",
        company_b_email="b@example.com",
        telegram_username="integration_user",
    )

    assert registration.telegram_user_id == 1001
    assert await repository.get_registration(1001) is not None
    assert await _outbox_row(database, 1001) == (1001, 0, None)


@pytest.mark.asyncio
async def test_registration_transaction_rolls_back_when_outbox_insert_fails(database: Database):
    repository = SQLiteCommandRepository(database)

    await database.connection.execute(
        """
        CREATE TRIGGER fail_lottery_outbox_insert
        BEFORE INSERT ON lottery_export_outbox
        BEGIN
            SELECT RAISE(ABORT, 'forced outbox failure');
        END;
        """
    )
    await database.connection.commit()

    with pytest.raises(RepositoryError, match="transaction failed"):
        await repository.create_registration(
            telegram_user_id=1002,
            full_name="Rollback User",
            company_a_email="rollback-a@example.com",
            company_b_email="rollback-b@example.com",
            telegram_username="rollback_user",
        )

    assert await repository.get_registration(1002) is None
    assert await _outbox_row(database, 1002) is None


class FailingExporter:
    def upsert(self, job):
        raise OSError("disk unavailable")


class RecordingExporter:
    def __init__(self):
        self.jobs = []

    def upsert(self, job):
        self.jobs.append(job)


@pytest.mark.asyncio
async def test_excel_failure_keeps_committed_registration_and_pending_outbox(database: Database):
    repository = SQLiteCommandRepository(database)
    await repository.create_registration(
        telegram_user_id=1003,
        full_name="Export Failure User",
        company_a_email="failure-a@example.com",
        company_b_email="failure-b@example.com",
        telegram_username="export_failure",
    )

    worker = LotteryExportWorker(repository, FailingExporter())
    assert await worker.process_pending_once() == 1

    assert await repository.get_registration(1003) is not None
    assert await _outbox_row(database, 1003) == (1003, 1, "disk unavailable")


@pytest.mark.asyncio
async def test_pending_outbox_is_processed_and_acknowledged(database: Database):
    repository = SQLiteCommandRepository(database)
    await repository.create_registration(
        telegram_user_id=1004,
        full_name="Pending User",
        company_a_email="pending-a@example.com",
        company_b_email="pending-b@example.com",
        telegram_username="pending_user",
    )

    exporter = RecordingExporter()
    worker = LotteryExportWorker(repository, exporter)

    assert await worker.process_pending_once() == 1
    assert [job.telegram_user_id for job in exporter.jobs] == [1004]
    assert await _outbox_row(database, 1004) is None
    assert await repository.get_registration(1004) is not None


@pytest.mark.asyncio
async def test_pending_outbox_survives_database_restart(tmp_path: Path):
    path = tmp_path / "restart.db"

    first_database = Database(str(path))
    await first_database.connect()
    try:
        repository = SQLiteCommandRepository(first_database)
        await repository.create_registration(
            telegram_user_id=1005,
            full_name="Restart User",
            company_a_email="restart-a@example.com",
            company_b_email="restart-b@example.com",
            telegram_username="restart_user",
        )
    finally:
        await first_database.close()

    second_database = Database(str(path))
    await second_database.connect()
    try:
        repository = SQLiteCommandRepository(second_database)
        pending = await repository.get_pending_exports()
        assert [job.telegram_user_id for job in pending] == [1005]
    finally:
        await second_database.close()


@pytest.mark.asyncio
async def test_real_excel_export_is_idempotent_for_replayed_outbox(tmp_path: Path, database: Database):
    repository = SQLiteCommandRepository(database)
    await repository.create_registration(
        telegram_user_id=1006,
        full_name="Idempotent User",
        company_a_email="idempotent-a@example.com",
        company_b_email="idempotent-b@example.com",
        telegram_username="idempotent_user",
    )

    exporter = LotteryExcelExporter(str(tmp_path / "lottery.xlsx"))
    worker = LotteryExportWorker(repository, exporter)

    assert await worker.process_pending_once() == 1
    assert await _outbox_row(database, 1006) is None

    await repository.enqueue_registration(1006)
    assert await worker.process_pending_once() == 1

    workbook = None
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(tmp_path / "lottery.xlsx")
        sheet = workbook.active
        matching_rows = [
            row for row in sheet.iter_rows(min_row=2, values_only=True)
            if row[0] == 1006
        ]
        assert len(matching_rows) == 1
    finally:
        if workbook is not None:
            workbook.close()
